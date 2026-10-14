"""Banners (Pillow), a ready-to-host landing page, and export bundles."""
from __future__ import annotations

import html
import io
import json
import zipfile
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONTS = Path(__file__).parent / "fonts"
NAVY, GOLD, GOLD_INK, IVORY, MUTED, WHITE = (19, 41, 75), (201, 151, 58), (122, 90, 22), (251, 248, 242), (102, 112, 133), (255, 255, 255)
SIZES = {"feed": (1080, 1080), "story": (1080, 1920), "link": (1200, 628)}


def _font(kind: str, size: int, weight: int = 600):
    path = {"serif": "Fraunces.ttf", "italic": "Fraunces-Italic.ttf", "sans": "Manrope.ttf"}[kind]
    try:
        f = ImageFont.truetype(str(FONTS / path), size)
        try:
            axes = f.get_variation_axes()
            vals = []
            for a in axes:
                n = (a.get("name") or b"").decode() if isinstance(a.get("name"), bytes) else str(a.get("name"))
                if "eight" in n or n == "wght":
                    vals.append(weight)
                elif "ptical" in n or n == "opsz":
                    vals.append(min(144, max(9, size // 2)))
                else:
                    vals.append(a.get("default", 0))
            f.set_variation_by_axes(vals)
        except Exception:
            pass
        return f
    except Exception:
        return ImageFont.load_default()


def nice_date(iso: str, weekday: bool = True) -> str:
    d = date.fromisoformat(iso)
    return (d.strftime('%a, %b ') if weekday else d.strftime('%b ')) + str(d.day)


def _wrap(d: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    lines, cur = [], ""
    for w in text.split():
        t = f"{cur} {w}".strip()
        if d.textlength(t, font=font) > width and cur:
            lines.append(cur)
            cur = w
        else:
            cur = t
    if cur:
        lines.append(cur)
    return lines


def banner(kind: str, eyebrow: str, headline: str, detail: str, cta: str) -> bytes:
    W, H = SIZES[kind]
    im = Image.new("RGB", (W, H), IVORY)
    d = ImageDraw.Draw(im)
    # court-line motif
    cx, cy = int(W * 0.86), int(H * 0.16)
    for r in (int(min(W, H) * .28), int(min(W, H) * .20), int(min(W, H) * .12)):
        d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=(232, 214, 176), width=3)
    d.rectangle([0, H - 16, W, H], fill=GOLD)
    pad = int(W * 0.075)
    wide = W / H > 1.5
    base = H / 628 if wide else W / 1080
    y = int(H * (0.30 if kind == 'story' else 0.16 if wide else 0.14))
    if kind == 'story':
        base = 1.25
    d.text((pad, y), eyebrow.upper()[:60], font=_font("sans", int(24 * base), 800), fill=GOLD_INK)
    y += int(56 * base)
    size = int((64 if wide else 96 if kind == "feed" else 110) * base)
    maxw = int(W * (0.62 if wide else 0.84))
    f = _font("serif", size, 600)
    lines = _wrap(d, headline, f, maxw)
    while len(lines) > (2 if wide else 4) and size > 36:
        size -= 6
        f = _font("serif", size, 600)
        lines = _wrap(d, headline, f, maxw)
    for ln in lines:
        d.text((pad, y), ln, font=f, fill=NAVY)
        y += int(size * 1.08)
    y += int(18 * base)
    fd = _font("sans", int((24 if wide else 34) * base), 500)
    for ln in _wrap(d, detail, fd, maxw)[:3]:
        d.text((pad, y), ln, font=fd, fill=MUTED)
        y += int((34 if wide else 46) * base)
    fc = _font("sans", int((24 if wide else 34) * base), 800)
    tw = d.textlength(cta, font=fc)
    bh = int((60 if wide else 86) * base)
    bx, by = (W - pad - tw - 70 * base, H // 2 - bh // 2) if wide else (pad, H - int(H * (0.12 if kind == "story" else 0.16)) - bh)
    d.rounded_rectangle([bx, by, bx + tw + 70 * base, by + bh], radius=bh // 2, fill=NAVY)
    d.text((bx + 35 * base, by + bh / 2), cta, font=fc, fill=WHITE, anchor="lm")
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def banners_for(kit: dict) -> dict[str, bytes]:
    inp, copy = kit["input"], kit["copy"]
    eyebrow = f"New campus · {inp['city']}"
    headline = copy["positioning"]["headline"]
    first_oh = next((t for t in kit["timeline"] if t["task"] == "Open House #1"), None)
    detail = f"Free Open House {nice_date(first_oh['due'])}" if first_oh else copy["positioning"]["tagline"]
    return {k: banner(k, eyebrow, headline, detail, "Reserve your spot  →") for k in SIZES}


def landing_page(kit: dict) -> str:
    inp, c = kit["input"], kit["copy"]
    e = html.escape
    events = [t for t in kit["timeline"] if t["category"] == "Events" and t["task"] != "Family orientation night"]
    ev_html = "".join(f'<li><b>{e(nice_date(t["due"]))}</b><span>{e(t["task"])}</span></li>' for t in events)
    pillars = "".join(f'<div class="p"><h3>{e(p["title"])}</h3><p>{e(p["body"])}</p></div>' for p in c["website"]["pillars"])
    faq = "".join(f'<details><summary>{e(q["q"])}</summary><p>{e(q["a"])}</p></details>' for q in c["website"]["faq"])
    pts = "".join(f"<li>{e(x)}</li>" for x in c["positioning"]["selling_points"])
    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>{e(kit['name'])}</title><meta name="description" content="{e(c['positioning']['tagline'])}">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Fraunces:ital,opsz,wght@0,9..144,600;1,9..144,500&family=Plus+Jakarta+Sans:wght@400;600;700&display=swap" rel="stylesheet">
<style>:root{{--navy:#13294B;--gold:#C9973A;--ink:#14213D;--muted:#667085;--bg:#F8F6F1;--panel:#fff;--line:#ECE8DF}}
@media (prefers-color-scheme:dark){{:root{{--ink:#EEF1F6;--muted:#9AA4B5;--bg:#0F1522;--panel:#161E2E;--line:#253045;--navy:#E7C77E}}}}
*{{box-sizing:border-box}}body{{margin:0;font-family:"Plus Jakarta Sans",system-ui,Arial,sans-serif;color:var(--ink);background:var(--bg);line-height:1.6}}
.w{{width:min(1080px,100% - 40px);margin:0 auto}}h1,h2,h3{{font-family:Fraunces,Georgia,serif;letter-spacing:-.01em}}
header{{padding:72px 0 56px;background:radial-gradient(800px 360px at 90% -10%,rgba(201,151,58,.14),transparent 60%)}}
.eb{{font-size:.75rem;letter-spacing:.18em;text-transform:uppercase;color:var(--gold);font-weight:700}}
h1{{font-size:clamp(2.4rem,6vw,4rem);line-height:1.04;margin:12px 0 16px;max-width:15ch}}header p{{font-size:1.15rem;color:var(--muted);max-width:40em}}
.btn{{display:inline-block;margin-top:22px;background:var(--navy);color:#fff;padding:14px 26px;border-radius:12px;font-weight:700;text-decoration:none}}
section{{padding:56px 0;border-top:1px solid var(--line)}}ul.pts{{padding-left:20px}}ul.pts li{{margin-bottom:8px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:16px}}.p{{background:var(--panel);border:1px solid var(--line);border-radius:16px;padding:22px}}
.p h3{{margin:0 0 6px}}.p p{{margin:0;color:var(--muted)}}ol.ev{{list-style:none;padding:0;margin:0}}ol.ev li{{display:flex;gap:18px;padding:14px 0;border-bottom:1px solid var(--line)}}ol.ev b{{min-width:120px}}
details{{border-bottom:1px solid var(--line);padding:16px 0}}summary{{cursor:pointer;font-weight:700}}details p{{color:var(--muted)}}
footer{{padding:36px 0;color:var(--muted);font-size:.85rem;border-top:1px solid var(--line)}}</style></head><body>
<header><div class="w"><span class="eb">{e(inp['city'])} · Opening {e(date.fromisoformat(inp['opening']).strftime('%B %Y'))}</span><h1>{e(c['website']['hero_headline'])}</h1><p>{e(c['website']['hero_sub'])}</p><a class="btn" href="#events">See upcoming events</a></div></header>
<section><div class="w"><h2>{e(c['positioning']['headline'])}</h2><ul class="pts">{pts}</ul></div></section>
<section><div class="w"><h2>The program</h2><div class="grid">{pillars}</div></div></section>
<section id="events"><div class="w"><h2>Upcoming events</h2><ol class="ev">{ev_html}</ol></div></section>
<section><div class="w"><h2>Questions</h2>{faq}</div></section>
<footer><div class="w">Concept launch page generated by AI Campus Launch Copilot (Arbaz Abdul). Not an official Texas Sports Academy website.</div></footer></body></html>"""


def checklist_xlsx(kit: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Launch checklist"
    ws.append([kit["name"]])
    ws["A1"].font = Font(bold=True, size=14, color="13294B")
    ws.append([f"Opening {kit['input']['opening']} · Target {kit['input']['target_students']} students"])
    ws.append([])
    hdr = ["Due date", "Task", "Owner", "Category", "Done"]
    ws.append(hdr)
    for i in range(1, 6):
        c = ws.cell(row=4, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="13294B")
    for t in kit["timeline"]:
        ws.append([date.fromisoformat(t["due"]), t["task"], t["owner"], t["category"], "Yes" if t.get("done") else "No"])
        ws.cell(row=ws.max_row, column=1).number_format = "mmm d, yyyy"
    for col, w in zip("ABCDE", (14, 48, 18, 14, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    s2 = wb.create_sheet("Hiring plan")
    s2.append(["Role", "Team", "Openings"])
    for r in kit["staffing"]:
        s2.append([r["title"], r["team"], r["openings"]])
    s3 = wb.create_sheet("Ad budget")
    s3.append(["Phase", "Budget (USD)"])
    for r in kit["budget"]:
        s3.append([r["phase"], r["budget"]])
    s3.append(["Total", f"=SUM(B2:B{len(kit['budget']) + 1})"])
    for s in (s2, s3):
        for i, w in enumerate((44, 16, 12), start=1):
            s.column_dimensions[get_column_letter(i)].width = w
        for c in s[1]:
            c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def kit_markdown(kit: dict) -> str:
    c, i = kit["copy"], kit["input"]
    L = [f"# {kit['name']}", f"Opening {i['opening']} · {i['sport']} · Grades {i['grade_min']}–{i['grade_max']} · Target {i['target_students']} students",
         f"_Copy source: {kit['source']}_", "", "## Positioning", f"**{c['positioning']['headline']}**", c["positioning"]["tagline"], ""]
    L += [f"- {x}" for x in c["positioning"]["selling_points"]]
    L += ["", "## Launch announcement", c["launch_announcement"], "", "## Facebook and Instagram ads", f"Audience: {c['audience']['summary']} Ages {c['audience']['age_range']}. Interests: {', '.join(c['audience']['interests'])}.", ""]
    for a in c["ads"]:
        L += [f"### {a['name']}", f"- Primary text: {a['primary_text']}", f"- Headline: {a['headline']}", f"- Button: {a['cta']}", ""]
    L += ["## Ad budget"] + [f"- {b['phase']}: ${b['budget']:,}" for b in kit["budget"]]
    L += ["", "## Hiring plan"] + [f"- {r['title']} ({r['openings']}): {c['job_blurbs'].get(r['key'], '')}" for r in kit["staffing"]]
    L += ["", "## Launch checklist"] + [f"- [{'x' if t.get('done') else ' '}] {t['due']} · {t['task']} ({t['owner']})" for t in kit["timeline"]]
    return "\n".join(L) + "\n"


def bundle_zip(kit: dict) -> bytes:
    buf = io.BytesIO()
    slug = "".join(ch if ch.isalnum() else "-" for ch in kit["name"].lower()).strip("-")
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{slug}/launch-kit.md", kit_markdown(kit))
        z.writestr(f"{slug}/landing-page/index.html", landing_page(kit))
        z.writestr(f"{slug}/launch-checklist.xlsx", checklist_xlsx(kit))
        z.writestr(f"{slug}/kit.json", json.dumps(kit, indent=2))
        for k, b in banners_for(kit).items():
            z.writestr(f"{slug}/banners/{k}-{'x'.join(map(str, SIZES[k]))}.png", b)
    return buf.getvalue()
